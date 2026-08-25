from storage import load_data
from auth import load_users
from utils import farsi

TRANSACTIONS_FILE = "data/transactions.json"

def admin_report_summary():
    transactions = load_data(TRANSACTIONS_FILE)
    users = load_users()

    total_income = sum(t.get('amount', 0) for t in transactions if t.get('type') == 'income')
    total_expense = sum(t.get('amount', 0) for t in transactions if t.get('type') == 'expense')
    balance = total_income - total_expense

    total_count = len(transactions)
    user_count = len(users)

    print(farsi("\n=== گزارش کلی سیستم ==="))
    print(farsi(f"تعداد کل کاربران: {user_count}"))
    print(farsi(f"تعداد کل تراکنش‌ها: {total_count}"))
    print(farsi(f"مجموع درآمدها: {total_income:,.0f} تومان"))
    print(farsi(f"مجموع مخارج: {total_expense:,.0f} تومان"))
    print(farsi(f"تراز مالی کل (مانده): {balance:,.0f} تومان"))
    print(farsi("------------------------------"))
from storage import load_data
from utils import farsi

def show_user_summary(username):
    transactions = load_data("data/transactions.json")
    user_txs = [t for t in transactions if t.get("username") == username]
    
    if not user_txs:
        print(farsi("تراکنشی برای گزارش‌گیری یافت نشد."))
        return

    income = sum(t['amount'] for t in user_txs if t['type'] == 'income')
    expense = sum(t['amount'] for t in user_txs if t['type'] == 'expense')
    
    # گزارش دسته‌بندی
    print(farsi("\n--- گزارش تفکیکی مخارج ---"))
    cat_sums = {}
    for t in user_txs:
        if t['type'] == 'expense':
            cat = t['category']
            cat_sums[cat] = cat_sums.get(cat, 0) + t['amount']
    
    for cat, total in cat_sums.items():
        print(farsi(f"{cat}: {total:,.0f} تومان"))
    
    print(farsi("--------------------------"))
    print(farsi(f"مجموع درآمد: {income:,.0f}"))
    print(farsi(f"مجموع مخارج: {expense:,.0f}"))
    print(farsi(f"مانده حساب: {(income - expense):,.0f}"))
import matplotlib.pyplot as plt
from storage import load_data
from utils import farsi

def show_expense_chart(username):
    transactions = load_data("data/transactions.json")
    user_txs = [t for t in transactions if t.get("username") == username and t.get("type") == "expense"]

    if not user_txs:
        print(farsi("هیچ هزینه‌ای برای رسم نمودار ثبت نشده است."))
        return

    # جمع مبالغ به تفکیک دسته‌بندی
    category_totals = {}
    for t in user_txs:
        cat = t.get("category", "سایر")
        category_totals[cat] = category_totals.get(cat, 0) + t.get("amount", 0)

    # اصلاح جهت و اتصال حروف فارسی برای نمایش در نمودار
    categories = [farsi(cat) for cat in category_totals.keys()]
    amounts = list(category_totals.values())

    # پالت رنگی جذاب
    colors = plt.get_cmap('tab10').colors

    # رسم نمودار دونات (حلقه‌ای) مدرن
    plt.figure(figsize=(8, 6))
    wedges, texts, autotexts = plt.pie(
        amounts,
        labels=None,
        autopct='%1.1f%%',
        startangle=140,
        colors=colors,
        pctdistance=0.75,
        wedgeprops=dict(width=0.6, edgecolor='white', linewidth=2)
    )

    # خوانایی و بولد کردن درصدها
    for autotext in autotexts:
        autotext.set_fontsize(10)
        autotext.set_weight('bold')

    # راهنمای دسته‌بندی‌ها در کنار نمودار
    plt.legend(
        wedges,
        categories,
        title=farsi("دسته‌بندی‌ها"),
        loc="center left",
        bbox_to_anchor=(1, 0, 0.5, 1)
    )

    plt.title(f"Expense Breakdown - {username}", fontsize=13, pad=15)
    plt.tight_layout()

    print(farsi("در حال نمایش نمودار..."))
    plt.show()
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from storage import load_data
from utils import farsi

TRANSACTIONS_FILE = "data/transactions.json"

def export_transactions_csv(username=None):
    """
    خروجی CSV: اگر username داده شود فقط تراکنش‌های آن کاربر، وگرنه تمام تراکنش‌ها (برای ادمین)
    """
    transactions = load_data(TRANSACTIONS_FILE)
    if username:
        data = [t for t in transactions if t.get("username") == username]
        filename = f"transactions_{username}.csv"
        headers = ["شناسه", "نوع", "مبلغ (تومان)", "دسته‌بندی", "توضیحات", "تاریخ"]
    else:
        data = transactions
        filename = "transactions_all_admin.csv"
        headers = ["شناسه", "کاربر", "نوع", "مبلغ (تومان)", "دسته‌بندی", "توضیحات", "تاریخ"]

    if not data:
        print(farsi("تراکنشی برای ذخیره وجود ندارد."))
        return

    with open(filename, mode="w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        for t in data:
            t_type = "درآمد" if t.get("type") == "income" else "هزینه"
            row = [
                t.get("id", "-"),
                t.get("username") if not username else None,
                t_type,
                t.get("amount", 0),
                t.get("category", "-"),
                t.get("note", "-"),
                t.get("date", "-")
            ]
            if username:
                row.pop(1)
            writer.writerow(row)

    print(farsi(f"فایل CSV با موفقیت ایجاد شد: {filename}"))
def export_transactions_excel(username=None):
    """
    خروجی Excel با قالب‌بندی رنگی، هدرهای بولد و جهت راست‌به‌چپ
    """
    transactions = load_data(TRANSACTIONS_FILE)
    if username:
        data = [t for t in transactions if t.get("username") == username]
        filename = f"transactions_{username}.xlsx"
        headers = ["شناسه", "نوع", "مبلغ (تومان)", "دسته‌بندی", "توضیحات", "تاریخ"]
    else:
        data = transactions
        filename = "transactions_all_admin.xlsx"
        headers = ["شناسه", "کاربر", "نوع", "مبلغ (تومان)", "دسته‌بندی", "توضیحات", "تاریخ"]

    if not data:
        print(farsi("تراکنشی برای خروجی اکسل وجود ندارد."))
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.sheet_view.rightToLeft = True  # نمایش راست‌به‌چپ شیت

    # استایل هدر
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # درج داده‌ها
    for t in data:
        t_type = "درآمد" if t.get("type") == "income" else "هزینه"
        row = [
            t.get("id", "-"),
            t.get("username") if not username else None,
            t_type,
            t.get("amount", 0),
            t.get("category", "-"),
            t.get("note", "-"),
            t.get("date", "-")
        ]
        if username:
            row.pop(1)
        ws.append(row)

    # تنظیم خودکار عرض ستون‌ها
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    wb.save(filename)
    print(farsi(f"فایل Excel با موفقیت ذخیره شد: {filename}"))
